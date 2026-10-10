"""Formatos internos y comerciales, con una política única independiente de la interfaz."""
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile, ZIP_DEFLATED
import reportlab
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


POLICY = {
    'inventario': ('pdf', 'xlsx'), 'turnos': ('pdf', 'xlsx'),
    'servicios': ('pdf', 'xlsx', 'imagen', 'tarjetas'), 'productos': ('imagen', 'tarjetas'),
}
MIMES = {'pdf': 'application/pdf', 'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
         'png': 'image/png', 'zip': 'application/zip'}
FONT = Path(reportlab.__file__).parent / 'fonts' / 'Vera.ttf'


def validate(scope, fmt, comment):
    if scope not in POLICY or fmt not in POLICY[scope]:
        raise ValueError('Ese formato no está disponible para este contenido.')
    if not isinstance(comment, str) or len(comment) > 160:
        raise ValueError('El encabezado admite hasta 160 caracteres.')


def table(store, scope, query='', ids=None):
    if not isinstance(query, str) or len(query) > 120:
        raise ValueError('Búsqueda no válida.')
    if ids is not None and (not isinstance(ids, list) or len(ids) > 1000 or any(type(x) is not int or x <= 0 for x in ids)):
        raise ValueError('Selección no válida.')
    if hasattr(store, 'report'):
        columns, rows, items = store.report(scope, query, ids)
        if not items: raise ValueError('No hay registros para exportar.')
        if len(items) > 10000: raise ValueError('Reduce el informe a 10000 registros.')
        return columns, rows, items
    with store.connect() as db:
        db.execute('BEGIN')
        if scope == 'turnos':
            items = [dict(x) for x in db.execute('SELECT * FROM shifts ORDER BY id DESC LIMIT 10001')]
            columns = ['Turno', 'Apertura (UTC)', 'Cierre (UTC)', 'Fondo CUP', 'Esperado CUP', 'Contado CUP', 'Diferencia CUP']
            rows = [[x['id'], x['opened'], x['closed'] or 'Abierto', Decimal(x['fund']) / 100,
                     Decimal(x['expected']) / 100 if x['expected'] is not None else '',
                     Decimal(x['counted']) / 100 if x['counted'] is not None else '',
                     Decimal(x['counted'] - x['expected']) / 100 if x['closed'] else ''] for x in items]
        else:
            kind = 'servicio' if scope == 'servicios' else 'producto'
            items = []
            selection = set(ids) if ids is not None else None
            for row in db.execute('SELECT * FROM items WHERE kind=? ORDER BY name', (kind,)):
                if query.casefold() in row['name'].casefold() and (selection is None or row['id'] in selection):
                    items.append(dict(row))
                    if len(items) > 10000: raise ValueError('Reduce el informe a 10000 registros.')
            columns = ['Nombre', 'Precio CUP', 'Costo CUP'] + (['Existencias'] if kind == 'producto' else [])
            rows = [[x['name'], Decimal(x['price']) / 100, Decimal(x['cost']) / 100] +
                    ([x['stock']] if kind == 'producto' else []) for x in items]
        if len(items) > 10000:
            raise ValueError('Reduce el informe a 10000 registros.')
        if not items:
            raise ValueError('No hay registros para exportar.')
        return columns, rows, items


def xlsx(columns, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = 'SPVI'
    for row_index, values in enumerate([columns] + rows, 1):
        for col_index, value in enumerate(values, 1):
            cell = sheet.cell(row_index, col_index, value)
            # No permitir fórmulas ni hipervínculos ejecutables escritos en los nombres.
            if isinstance(value, str):
                cell.data_type = 's'
            if row_index == 1:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='185E6E')
            elif isinstance(value, (int, Decimal)):
                cell.font = Font(bold=True)
                if 'CUP' in columns[col_index - 1]:
                    cell.number_format = '#,##0.00 "CUP"'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for i, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(i)].width = min(70, max(12, max(len(str(r[i - 1])) for r in [columns] + rows) + 2))
    output = BytesIO(); book.save(output)
    return output.getvalue()


def pdf(title, columns, rows):
    output = BytesIO()
    size = landscape(A4) if len(columns) > 5 else A4
    doc = SimpleDocTemplate(output, pagesize=size, leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    cell = styles['BodyText']; cell.fontSize = 8; cell.leading = 11
    data = [[Paragraph(escape(str(v)), cell) for v in columns]]
    for row in rows:
        data.append([Paragraph(('<b>' + escape(str(v)) + '</b>') if isinstance(v, (int, Decimal)) else escape(str(v)), cell) for v in row])
    grid = Table(data, colWidths=[(size[0] - 56) / len(columns)] * len(columns), repeatRows=1)
    grid.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#dcecef')),
                             ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f6f7')]),
                             ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
    doc.build([Paragraph(escape(title), styles['Title']), Spacer(1, 14), grid])
    return output.getvalue()


def lines(text, font, width):
    """Ajusta incluso palabras largas sin recortar el encabezado."""
    result, current = [], ''
    for char in text:
        if char == '\n' or font.getlength(current + char) > width:
            result.append(current); current = ''
        if char != '\n': current += char
    result.append(current)
    return result


def images(items, fmt, comment):
    if len(items) > 300:
        raise ValueError('Selecciona como máximo 300 artículos por exportación de imagen.')
    font = ImageFont.truetype(str(FONT), 32)
    title_font = ImageFont.truetype(str(FONT), 46)
    output = []
    groups = [[x] for x in items] if fmt == 'tarjetas' else [items[i:i + 20] for i in range(0, len(items), 20)]
    for index, group in enumerate(groups, 1):
        content = []
        if fmt == 'tarjetas' and comment.strip():
            content += [(line, title_font, '#185e6e') for line in lines(comment.strip(), title_font, 970)]
            content.append(('', font, '#18353c'))
        for item in group:
            content += [(line, title_font if fmt == 'tarjetas' else font, '#18353c')
                        for line in lines(item['name'], title_font if fmt == 'tarjetas' else font, 970)]
            content.append((f"{Decimal(item['price']) / 100:,.2f} CUP", title_font, '#185e6e'))
            content.append(('', font, '#18353c'))
        photo=group[0].get('photo') if fmt=='tarjetas' else None
        height = max(500, 100 + sum(f.size + 16 for _, f, _ in content)+(520 if photo else 0))
        image = Image.new('RGB', (1080, height), '#f3f6f7'); draw = ImageDraw.Draw(image)
        y = 50
        if photo:
            with Image.open(BytesIO(photo)) as picture:
                picture.thumbnail((970,480));image.paste(picture,((1080-picture.width)//2,y));y+=520
        for text, used_font, color in content:
            draw.text((54, y), text, font=used_font, fill=color); y += used_font.size + 16
        buffer = BytesIO(); image.save(buffer, format='PNG'); image.close()
        output.append((f'SPVI_{index}.png', buffer.getvalue()))
    if len(output) == 1:
        return output[0][1], 'png'
    buffer = BytesIO()
    with ZipFile(buffer, 'w', ZIP_DEFLATED) as archive:
        for name, content in output: archive.writestr(name, content)
    return buffer.getvalue(), 'zip'


def export(store, scope, fmt, comment='', query='', ids=None):
    validate(scope, fmt, comment)
    columns, rows, items = table(store, scope, query, ids)
    if fmt == 'pdf': return pdf('SPVI · ' + scope.capitalize(), columns, rows), 'pdf'
    if fmt == 'xlsx': return xlsx(columns, rows), 'xlsx'
    return images(items, fmt, comment)
